"""XLSX import: parse a workbook into a list of receive-new-box requests and
apply them via `services.boxes.create_box`. Best-effort, with row-indexed
reasons for the skipped rows so the UI can show actionable errors.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.box_files import normalize_box_file_reference
from app.models.boxes import Box
from app.models.pallets import clean_optional_pallet_number
from app.models.requests import BoxRequestOrigin
from app.models.users import User
from app.models.warehouses import Warehouse
from app.schemas.requests import InboundBoxItem
from app.services.acl import can_access
from app.services.boxes import (
    BoxAccessError,
    BoxRuleError,
    create_box,
    normalize_box_number,
    restore_archived_box,
)
from app.services.lots import (
    LotRuleError,
    find_lot,
    normalize_lot_name,
    resolve_lot_names_for_use,
    validate_lot_name,
)
from app.services.requests import (
    RequestRuleError,
    create_completed_receipt,
    create_staged_receipt,
    merge_inbound_items,
    receipt_requires_review,
)

# Hard caps: we do this on the request thread so we want a worst-case bound.
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000

# Header alias map: lowercased sheet header -> canonical key.
_HEADER_ALIASES = {
    "box_number": "box_number",
    "box number": "box_number",
    "box": "box_number",
    "lot": "lot",
    "lot number": "lot",
    "lot_number": "lot",
    "pallet_number": "pallet_number",
    "pallet number": "pallet_number",
    "pallet": "pallet_number",
    "pallet_id": "pallet_id",
    "pallet id": "pallet_id",
    "contents": "contents",
    "file_reference": "file_reference",
    "file reference": "file_reference",
    "reference": "file_reference",
    "file_description": "file_description",
    "file description": "file_description",
    "warehouse_id": "warehouse_id",
    "warehouse id": "warehouse_id",
    "warehouse": "warehouse",
    "building": "warehouse",
}


@dataclass
class ImportSkipEntry:
    row: int
    box_number: str | None
    reason: str


@dataclass
class ImportOutcome:
    created: list[Box] = field(default_factory=list)
    restored: list[Box] = field(default_factory=list)
    skipped: list[ImportSkipEntry] = field(default_factory=list)
    receipt_request_ids: list[int] = field(default_factory=list)
    staged_receipt_ids: list[int] = field(default_factory=list)


def _commit_import_receipts(
    db: Session,
    *,
    user: User,
    outcome: ImportOutcome,
) -> None:
    imported_boxes = outcome.created + outcome.restored
    if not imported_boxes:
        return
    boxes_by_warehouse: dict[int, list[Box]] = {}
    for box in imported_boxes:
        boxes_by_warehouse.setdefault(box.current_warehouse_id, []).append(box)
    for warehouse_id, boxes in boxes_by_warehouse.items():
        receipt = create_completed_receipt(
            db,
            user=user,
            warehouse_id=warehouse_id,
            boxes=boxes,
            origin=BoxRequestOrigin.xlsx_import,
            note="Completed automatically from XLSX box import.",
            commit=False,
        )
        outcome.receipt_request_ids.append(receipt.id)
    db.commit()
    for box in imported_boxes:
        db.refresh(box)


def import_mapped_boxes(
    db: Session,
    *,
    user: User,
    warehouse_id: int,
    items: list[InboundBoxItem | tuple],
    restore_archived: bool = False,
) -> ImportOutcome:
    outcome = ImportOutcome()
    warehouse = db.get(Warehouse, warehouse_id)
    if warehouse is None or not warehouse.is_active:
        raise BoxRuleError(f"warehouse {warehouse_id} does not exist or is archived")
    if not can_access(user, warehouse_id):
        raise BoxAccessError(f"no access to warehouse {warehouse_id}")
    canonical_rows: list[InboundBoxItem] = []
    for position, raw_item in enumerate(items, start=1):
        if isinstance(raw_item, InboundBoxItem):
            box_number = raw_item.box_number
            lot = raw_item.lot
            pallet_number = raw_item.pallet_number
            pallet_id = raw_item.pallet_id
            contents = raw_item.contents
            files = raw_item.files
        else:
            box_number, lot, pallet_number, pallet_id, contents, *remainder = raw_item
            files = remainder[0] if remainder else None
        try:
            cleaned_number = normalize_box_number(box_number)
            cleaned_lot = validate_lot_name(lot)
        except (BoxRuleError, LotRuleError) as exc:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=position,
                    box_number=box_number,
                    reason=str(exc),
                )
            )
            continue
        pallet_number = clean_optional_pallet_number(pallet_number)
        if pallet_id is not None and pallet_number is None:
            raise BoxRuleError("pallet_id requires pallet_number")
        canonical_rows.append(
            InboundBoxItem(
                box_number=cleaned_number,
                lot=cleaned_lot,
                pallet_number=pallet_number,
                pallet_id=pallet_id,
                contents=contents,
                files=files,
            )
        )
    try:
        merged_rows = merge_inbound_items(canonical_rows)
    except RequestRuleError as exc:
        raise BoxRuleError(str(exc)) from exc
    canonical_items = [
        (position, item)
        for position, item in enumerate(merged_rows, start=1)
    ]
    if not canonical_items:
        return outcome
    if receipt_requires_review(
        warehouse,
        BoxRequestOrigin.xlsx_import,
        quantity=len(canonical_items),
    ):
        receipt = create_staged_receipt(
            db,
            user=user,
            warehouse_id=warehouse_id,
            items=[
                item for _position, item in canonical_items
            ],
            origin=BoxRequestOrigin.xlsx_import,
            restore_archived=restore_archived,
        )
        outcome.staged_receipt_ids.append(receipt.id)
        return outcome
    try:
        lots_by_name = resolve_lot_names_for_use(
            db,
            user=user,
            names=[item.lot for _position, item in canonical_items],
            warehouse_id=warehouse_id,
        )
    except LotRuleError as exc:
        raise BoxRuleError(str(exc)) from exc
    for position, item in canonical_items:
        lot_record = lots_by_name[normalize_lot_name(item.lot)]
        try:
            restored = (
                restore_archived_box(
                    db,
                    user=user,
                    box_number=item.box_number,
                    lot_id=lot_record.id,
                    pallet_number=item.pallet_number,
                    pallet_id=item.pallet_id,
                    contents=item.contents,
                    files=item.files,
                    warehouse_id=warehouse_id,
                    note="Explicitly restored during mapped XLSX import.",
                    commit=False,
                )
                if restore_archived
                else None
            )
            box = restored or create_box(
                db,
                user=user,
                box_number=item.box_number,
                lot_id=lot_record.id,
                pallet_number=item.pallet_number,
                pallet_id=item.pallet_id,
                contents=item.contents,
                files=item.files,
                warehouse_id=warehouse_id,
                commit=False,
            )
        except BoxRuleError as exc:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=position,
                    box_number=item.box_number,
                    reason=str(exc),
                )
            )
            continue
        if restored is not None:
            outcome.restored.append(box)
        else:
            outcome.created.append(box)
    _commit_import_receipts(db, user=user, outcome=outcome)
    return outcome


def _normalise_headers(raw: list[object | None]) -> dict[str, int]:
    out: dict[str, int] = {}
    for idx, value in enumerate(raw):
        if value is None:
            continue
        key = str(value).strip().lower()
        canonical = _HEADER_ALIASES.get(key)
        if canonical and canonical not in out:
            out[canonical] = idx
    return out


def _cell_str(cell: object | None) -> str:
    if cell is None:
        return ""
    if isinstance(cell, float) and cell.is_integer():
        return str(int(cell))
    return str(cell).strip()


def _mapped_cell(
    headers: dict[str, int], cells: list[object | None], name: str
) -> str:
    idx = headers.get(name)
    if idx is None or idx >= len(cells):
        return ""
    return _cell_str(cells[idx])


def import_boxes_xlsx(
    db: Session,
    *,
    user: User,
    file_bytes: bytes,
    default_warehouse_id: int | None = None,
    restore_archived: bool = False,
) -> ImportOutcome:
    if len(file_bytes) > MAX_FILE_BYTES:
        raise BoxRuleError(
            f"file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB"
        )
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises a few different types
        raise BoxRuleError(f"could not read workbook: {exc}") from exc

    ws = wb.active
    if ws is None:
        raise BoxRuleError("workbook has no sheets")

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration as exc:
        raise BoxRuleError("workbook is empty") from exc

    headers = _normalise_headers(list(header_row))
    if "box_number" not in headers:
        raise BoxRuleError(
            "missing required column 'box_number' "
            "(recognised headers: box_number, lot, file_reference, contents, "
            "warehouse_id, warehouse)"
        )
    if "lot" not in headers:
        raise BoxRuleError(
            "missing required column 'lot' "
            "(recognised headers: box_number, lot, file_reference, contents, "
            "warehouse_id, warehouse)"
        )
    # Pre-resolve warehouses so we don't hit the DB once per row.
    warehouses_by_id = {
        w.id: w
        for w in db.scalars(
            select(Warehouse).where(Warehouse.is_active.is_(True))
        ).all()
        if can_access(user, w.id)
    }
    warehouses_by_name = {
        w.name.strip().lower(): w for w in warehouses_by_id.values()
    }
    if default_warehouse_id is not None and default_warehouse_id not in warehouses_by_id:
        raise BoxRuleError(
            f"default warehouse {default_warehouse_id} does not exist or is archived"
        )

    outcome = ImportOutcome()
    pending_by_warehouse: dict[
        int, list[tuple[int, InboundBoxItem]]
    ] = {}
    pending_identity: dict[tuple[str, str], tuple[int, int]] = {}
    pending_file_targets: dict[tuple[str, str], str] = {}
    blocked_identities: dict[tuple[str, str], str] = {}

    for offset, row in enumerate(rows_iter, start=2):  # row 1 was the header
        if offset - 1 > MAX_ROWS:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=offset,
                    box_number=None,
                    reason=f"row limit {MAX_ROWS} exceeded; remainder ignored",
                )
            )
            break

        if row is None or all(cell is None or _cell_str(cell) == "" for cell in row):
            continue  # blank row

        cells = list(row)

        raw_box_number = _mapped_cell(headers, cells, "box_number")
        lot = _mapped_cell(headers, cells, "lot")
        pallet_number = clean_optional_pallet_number(
            _mapped_cell(headers, cells, "pallet_number")
        )
        pallet_id_raw = _mapped_cell(headers, cells, "pallet_id")
        contents_raw = _mapped_cell(headers, cells, "contents")
        contents = contents_raw or None
        file_reference = _mapped_cell(headers, cells, "file_reference")
        file_description = _mapped_cell(headers, cells, "file_description") or None
        wh_id_raw = _mapped_cell(headers, cells, "warehouse_id")
        wh_name_raw = _mapped_cell(headers, cells, "warehouse")

        if not raw_box_number:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=offset, box_number=None, reason="box_number is empty"
                )
            )
            continue
        # Apply the same canonical-form rule as the API: numeric only,
        # zero-padded to 3 digits. Non-numeric rows are surfaced with the
        # row index so the operator can fix the sheet.
        try:
            box_number = normalize_box_number(raw_box_number)
        except BoxRuleError as exc:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=offset, box_number=raw_box_number, reason=str(exc)
                )
            )
            continue
        if not lot:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=offset, box_number=box_number, reason="lot is empty"
                )
            )
            continue
        try:
            lot = validate_lot_name(lot)
        except LotRuleError as exc:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=offset,
                    box_number=box_number,
                    reason=str(exc),
                )
            )
            continue
        pallet_id: int | None = None
        if pallet_id_raw:
            try:
                pallet_id = int(pallet_id_raw)
            except ValueError:
                outcome.skipped.append(
                    ImportSkipEntry(
                        row=offset,
                        box_number=box_number,
                        reason=f"pallet_id {pallet_id_raw!r} is not an integer",
                    )
                )
                continue
            if pallet_id < 1:
                outcome.skipped.append(
                    ImportSkipEntry(
                        row=offset,
                        box_number=box_number,
                        reason="pallet_id must be positive",
                    )
                )
                continue
        if pallet_id is not None and pallet_number is None:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=offset,
                    box_number=box_number,
                    reason="pallet_id requires pallet_number",
                )
            )
            continue
        if "file_reference" in headers and not file_reference:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=offset,
                    box_number=box_number,
                    reason="file_reference is required for every mapped file row",
                )
            )
            continue

        warehouse_id: int | None = None
        if wh_id_raw:
            try:
                warehouse_id = int(wh_id_raw)
            except ValueError:
                outcome.skipped.append(
                    ImportSkipEntry(
                        row=offset,
                        box_number=box_number,
                        reason=f"warehouse_id {wh_id_raw!r} is not an integer",
                    )
                )
                continue
            if warehouse_id not in warehouses_by_id:
                outcome.skipped.append(
                    ImportSkipEntry(
                        row=offset,
                        box_number=box_number,
                        reason=f"warehouse {warehouse_id} does not exist",
                    )
                )
                continue
        elif wh_name_raw:
            wh = warehouses_by_name.get(wh_name_raw.strip().lower())
            if wh is None:
                outcome.skipped.append(
                    ImportSkipEntry(
                        row=offset,
                        box_number=box_number,
                        reason=f"unknown warehouse name {wh_name_raw!r}",
                    )
                )
                continue
            warehouse_id = wh.id
        elif default_warehouse_id is not None:
            warehouse_id = default_warehouse_id
        else:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=offset,
                    box_number=box_number,
                    reason=(
                        "no warehouse for this row "
                        "(set warehouse_id/warehouse column or pass a default)"
                    ),
                )
            )
            continue

        item = InboundBoxItem(
            box_number=box_number,
            lot=lot,
            pallet_number=pallet_number,
            pallet_id=pallet_id,
            contents=contents,
            files=(
                [
                    {
                        "reference": file_reference,
                        "description": file_description,
                    }
                ]
                if file_reference
                else None
            ),
        )
        pair = (normalize_lot_name(lot), box_number)
        for file in item.files or []:
            file_identity = (
                pair[0],
                normalize_box_file_reference(file.reference),
            )
            previous_box = pending_file_targets.get(file_identity)
            if previous_box is not None and previous_box != box_number:
                raise BoxRuleError(
                    f"file reference {file.reference!r} targets two boxes "
                    f"in lot {lot!r}"
                )
            pending_file_targets[file_identity] = box_number
        blocked_reason = blocked_identities.get(pair)
        if blocked_reason is not None:
            outcome.skipped.append(
                ImportSkipEntry(
                    row=offset,
                    box_number=box_number,
                    reason=blocked_reason,
                )
            )
            continue
        duplicate = pending_identity.get(pair)
        if duplicate is not None:
            existing_warehouse_id, existing_index = duplicate
            if existing_warehouse_id != warehouse_id:
                reason = (
                    "duplicate (lot, box_number) rows specify different warehouses"
                )
                blocked_identities[pair] = reason
                existing_offset, _existing_item = pending_by_warehouse[
                    existing_warehouse_id
                ][existing_index]
                outcome.skipped.extend(
                    (
                        ImportSkipEntry(
                            row=existing_offset,
                            box_number=box_number,
                            reason=reason,
                        ),
                        ImportSkipEntry(
                            row=offset,
                            box_number=box_number,
                            reason=reason,
                        ),
                    )
                )
                continue
            existing_offset, existing_item = pending_by_warehouse[warehouse_id][
                existing_index
            ]
            try:
                merged = merge_inbound_items([existing_item, item])[0]
            except RequestRuleError as exc:
                reason = str(exc)
                blocked_identities[pair] = reason
                outcome.skipped.extend(
                    (
                        ImportSkipEntry(
                            row=existing_offset,
                            box_number=box_number,
                            reason=reason,
                        ),
                        ImportSkipEntry(
                            row=offset,
                            box_number=box_number,
                            reason=reason,
                        ),
                    )
                )
                continue
            pending_by_warehouse[warehouse_id][existing_index] = (
                existing_offset,
                merged,
            )
            continue
        pending_rows = pending_by_warehouse.setdefault(warehouse_id, [])
        pending_identity[pair] = (warehouse_id, len(pending_rows))
        pending_rows.append((offset, item))

    for warehouse_id, pending_rows in pending_by_warehouse.items():
        pending_rows = [
            (offset, item)
            for offset, item in pending_rows
            if (normalize_lot_name(item.lot), item.box_number)
            not in blocked_identities
        ]
        if not pending_rows:
            continue
        warehouse = warehouses_by_id[warehouse_id]
        if receipt_requires_review(
            warehouse,
            BoxRequestOrigin.xlsx_import,
            quantity=len(pending_rows),
        ):
            staged_items: list[InboundBoxItem] = []
            for offset, item in pending_rows:
                lot_record = find_lot(db, item.lot)
                existing = (
                    db.scalar(
                        select(Box).where(
                            Box.box_number == item.box_number,
                            Box.lot_id == lot_record.id,
                        )
                    )
                    if lot_record is not None
                    else None
                )
                if existing is not None and (
                    existing.archived_at is None or not restore_archived
                ):
                    reason = (
                        f"box {item.box_number!r} already exists in lot {item.lot!r}"
                        if existing.archived_at is None
                        else (
                            f"box {item.box_number!r} in lot {item.lot!r} is archived; "
                            "enable restore_archived to restore it"
                        )
                    )
                    outcome.skipped.append(
                        ImportSkipEntry(
                            row=offset,
                            box_number=item.box_number,
                            reason=reason,
                        )
                    )
                    continue
                staged_items.append(item)
            if staged_items:
                receipt = create_staged_receipt(
                    db,
                    user=user,
                    warehouse_id=warehouse_id,
                    items=staged_items,
                    origin=BoxRequestOrigin.xlsx_import,
                    restore_archived=restore_archived,
                )
                outcome.staged_receipt_ids.append(receipt.id)
            continue

        try:
            lots_by_name = resolve_lot_names_for_use(
                db,
                user=user,
                names=[item.lot for _offset, item in pending_rows],
                warehouse_id=warehouse_id,
            )
        except LotRuleError as exc:
            for offset, item in pending_rows:
                outcome.skipped.append(
                    ImportSkipEntry(
                        row=offset,
                        box_number=item.box_number,
                        reason=str(exc),
                    )
                )
            continue

        for offset, item in pending_rows:
            lot_record = lots_by_name[normalize_lot_name(item.lot)]
            try:
                restored = (
                    restore_archived_box(
                        db,
                        user=user,
                        box_number=item.box_number,
                        lot_id=lot_record.id,
                        pallet_number=item.pallet_number,
                        pallet_id=item.pallet_id,
                        contents=item.contents,
                        files=item.files,
                        warehouse_id=warehouse_id,
                        note=f"Explicitly restored from XLSX row {offset}.",
                        commit=False,
                    )
                    if restore_archived
                    else None
                )
                box = restored or create_box(
                    db,
                    user=user,
                    box_number=item.box_number,
                    lot_id=lot_record.id,
                    pallet_number=item.pallet_number,
                    pallet_id=item.pallet_id,
                    contents=item.contents,
                    files=item.files,
                    warehouse_id=warehouse_id,
                    commit=False,
                )
            except BoxRuleError as exc:
                outcome.skipped.append(
                    ImportSkipEntry(
                        row=offset,
                        box_number=item.box_number,
                        reason=str(exc),
                    )
                )
                continue
            if restored is not None:
                outcome.restored.append(box)
            else:
                outcome.created.append(box)

    _commit_import_receipts(db, user=user, outcome=outcome)
    return outcome
