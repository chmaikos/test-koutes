from __future__ import annotations

import io
from datetime import UTC, datetime

from openpyxl import load_workbook
from sqlalchemy import event, select

from app.deps import get_current_user
from app.events import bus
from app.main import app
from app.models import (
    Box,
    BoxStatus,
    Lot,
    LotEvent,
    LotEventType,
    ReceiptMode,
    UserRole,
    Warehouse,
)
from app.models.requests import BoxRequestOrigin
from app.schemas.requests import InboundBoxItem
from app.services.boxes import create_box as _create_box
from app.services.lots import list_lot_summaries
from app.services.requests import create_staged_receipt

create_box = _create_box


def _restrict_to(session, user, *warehouse_ids: int) -> None:
    user.warehouses = [
        session.get(Warehouse, warehouse_id) for warehouse_id in warehouse_ids
    ]
    session.commit()


def test_lot_metrics_completion_archives_and_zero_denominator(session, make_user):
    admin = make_user(UserRole.admin)
    lot = Lot(name="All Statuses")
    session.add(lot)
    session.flush()
    for index, status in enumerate(BoxStatus, start=1):
        create_box(
            session,
            user=admin,
            box_number=str(index),
            lot_id=lot.id,
            warehouse_id=1,
            initial_status=status,
        )
        if status == BoxStatus.received:
            archived = create_box(
                session,
                user=admin,
                box_number="99",
                lot_id=lot.id,
                warehouse_id=1,
            )
            archived.archived_at = datetime.now(UTC)
            session.commit()

    quarantined_only = create_box(
        session,
        user=admin,
        box_number="1",
        lot="No Eligible",
        warehouse_id=1,
        initial_status=BoxStatus.quarantined,
    )
    rows, total = list_lot_summaries(
        session, user=admin, sort_by="name", sort_dir="asc"
    )
    assert total == 2
    summary = next(row for row in rows if row.id == lot.id)
    assert summary.box_count == 6
    assert summary.physical_box_count == 5
    assert summary.status_counts == {status.value: 1 for status in BoxStatus}
    assert summary.eligible_box_count == 5
    assert summary.completed_box_count == 3
    assert summary.completion_percent == 60.0
    assert summary.progress_state == "in_progress"
    assert summary.warehouse_names == ["Building 1"]

    zero = next(row for row in rows if row.id == quarantined_only.lot_id)
    assert zero.eligible_box_count == 0
    assert zero.completion_percent is None
    assert zero.progress_state == "no_eligible"


def test_lot_acl_scopes_before_aggregation_and_includes_staged_receipts(
    session, make_user
):
    admin = make_user(UserRole.admin)
    restricted = make_user(UserRole.operator)
    lot = Lot(name="Shared Lot")
    session.add(lot)
    session.flush()
    create_box(
        session,
        user=admin,
        box_number="1",
        lot_id=lot.id,
        warehouse_id=1,
        initial_status=BoxStatus.incomplete,
    )
    create_box(
        session,
        user=admin,
        box_number="2",
        lot_id=lot.id,
        warehouse_id=2,
        initial_status=BoxStatus.received,
    )
    staged_receipt = create_staged_receipt(
        session,
        user=admin,
        warehouse_id=1,
        items=[
            InboundBoxItem(
                box_number="3",
                lot="  SHARED\tLOT ",
                pallet_number="PALLET-SHARED",
            )
        ],
        origin=BoxRequestOrigin.manual_entry,
    )
    assert staged_receipt.items[0].lot_id == lot.id
    assert staged_receipt.items[0].lot == "Shared Lot"
    hidden = create_box(
        session,
        user=admin,
        box_number="1",
        lot="Hidden Lot",
        warehouse_id=2,
    )
    _restrict_to(session, restricted, 1)

    rows, total = list_lot_summaries(session, user=restricted)
    assert total == 1
    assert rows[0].id == lot.id
    assert rows[0].box_count == 1
    assert rows[0].completion_percent == 100.0
    assert rows[0].warehouse_names == ["Building 1"]
    assert rows[0].staged_receipt_count == 1
    assert rows[0].acl_scoped is True
    assert hidden.lot_id not in {row.id for row in rows}


def test_mapped_import_returns_forbidden_outside_warehouse_acl(
    client, session, make_user
):
    restricted = make_user(UserRole.operator)
    _restrict_to(session, restricted, 1)
    app.dependency_overrides[get_current_user] = lambda: restricted

    response = client.post(
        "/api/boxes/import-mapped",
        json={
            "warehouse_id": 2,
            "items": [
                {
                    "box_number": "1",
                    "lot": "Hidden Import",
                    "pallet_number": "PALLET-HIDDEN",
                }
            ],
        },
    )

    assert response.status_code == 403
    assert session.scalar(select(Lot).where(Lot.name == "Hidden Import")) is None


def test_lot_api_search_sort_filter_pagination_exports_and_options(client):
    ids = []
    for name, status, warehouse_id in (
        ("Alpha Lot", BoxStatus.received, 1),
        ("Beta Lot", BoxStatus.incomplete, 1),
        ("Gamma Lot", BoxStatus.quarantined, 2),
    ):
        response = client.post(
            "/api/boxes",
            json={
                "box_number": "1",
                "lot": name,
                "pallet_number": f"PALLET-{name}",
                "warehouse_id": warehouse_id,
            },
        )
        assert response.status_code == 201
        box = response.json()
        ids.append(box["lot_id"])
        if status != BoxStatus.received:
            from app import db as db_module

            with db_module.SessionLocal() as db:
                stored = db.get(Box, box["id"])
                stored.status = status
                db.commit()

    searched = client.get(
        "/api/lots",
        params={
            "search": "lot",
            "warehouse_id": 1,
            "sort_by": "completion",
            "sort_dir": "desc",
            "page": 1,
            "page_size": 1,
        },
    )
    assert searched.status_code == 200, searched.text
    assert searched.json()["total"] == 2
    assert searched.json()["items"][0]["name"] == "Beta Lot"

    complete = client.get("/api/lots", params={"progress_state": "complete"})
    assert {item["name"] for item in complete.json()["items"]} == {"Beta Lot"}
    no_eligible = client.get(
        "/api/lots", params={"progress_state": "no_eligible"}
    )
    assert {item["name"] for item in no_eligible.json()["items"]} == {"Gamma Lot"}

    options = client.get("/api/lots/options", params={"search": "  ALPHA   LOT "})
    assert options.status_code == 200
    assert options.json()["items"][0]["id"] == ids[0]
    assert options.json()["items"][0]["exact_normalized_match"] is True

    csv_response = client.get(
        "/api/exports/lots.csv", params={"warehouse_id": 1}
    )
    assert csv_response.status_code == 200
    assert "Alpha Lot" in csv_response.text
    assert "Gamma Lot" not in csv_response.text
    assert "LOT-0000000000" in csv_response.text
    assert "Completion Percent" in csv_response.text

    xlsx_response = client.get("/api/exports/lots.xlsx")
    assert xlsx_response.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
    try:
        assert workbook.sheetnames == ["Lot summary"]
    finally:
        workbook.close()


def test_lot_create_rename_reassignment_audit_conflicts_and_sse(
    client, session, make_user, monkeypatch
):
    published: list[tuple[str, dict[str, object]]] = []

    async def capture(event_type: str, data: dict[str, object]) -> None:
        published.append((event_type, data))

    monkeypatch.setattr(bus, "publish", capture)
    created = client.post(
        "/api/lots",
        json={"name": "  Canonical\tName ", "warehouse_id": 1},
    )
    assert created.status_code == 201, created.text
    lot = created.json()
    assert lot["name"] == "Canonical Name"
    repeated = client.post(
        "/api/lots",
        json={"name": "canonical name", "warehouse_id": 1},
    )
    assert repeated.status_code == 200
    assert repeated.json()["id"] == lot["id"]

    operator = make_user(UserRole.operator)
    app.dependency_overrides[get_current_user] = lambda: operator
    denied = client.patch(
        f"/api/lots/{lot['id']}/rename",
        json={
            "new_name": "Denied",
            "reason": "not admin",
            "expected_version": lot["version"],
        },
    )
    assert denied.status_code == 403

    admin = make_user(UserRole.admin)
    app.dependency_overrides[get_current_user] = lambda: admin
    renamed = client.patch(
        f"/api/lots/{lot['id']}/rename",
        json={
            "new_name": "Renamed Lot",
            "reason": "Corrected supplier label",
            "expected_version": lot["version"],
        },
    )
    assert renamed.status_code == 200, renamed.text
    stale = client.patch(
        f"/api/lots/{lot['id']}/rename",
        json={
            "new_name": "Stale",
            "reason": "stale edit",
            "expected_version": lot["version"],
        },
    )
    assert stale.status_code == 409
    assert stale.json()["detail"]["code"] == "version_conflict"
    assert stale.json()["detail"]["current"]["name"] == "Renamed Lot"

    source = client.post(
        "/api/boxes",
        json={
            "box_number": "7",
            "lot": "Source",
            "pallet_number": "PALLET-SOURCE",
            "warehouse_id": 1,
        },
    ).json()
    target = client.post(
        "/api/boxes",
        json={
            "box_number": "8",
            "lot": "Target",
            "pallet_number": "PALLET-TARGET",
            "warehouse_id": 1,
        },
    ).json()
    source_lot = session.get(Lot, source["lot_id"])
    mixed = client.patch(
        f"/api/boxes/{source['id']}",
        json={
            "lot_id": target["lot_id"],
            "expected_version": source_lot.version,
            "note": "Lot changes use the dedicated audited operation",
        },
    )
    assert mixed.status_code == 422
    reassigned = client.post(
        f"/api/boxes/{source['id']}/reassign-lot",
        json={
            "lot_id": target["lot_id"],
            "reason": "Verified physical label",
            "expected_lot_version": source_lot.version,
        },
    )
    assert reassigned.status_code == 200, reassigned.text
    assert reassigned.json()["lot_id"] == target["lot_id"]
    assert session.scalar(
        select(LotEvent).where(
            LotEvent.lot_id == source["lot_id"],
            LotEvent.event_type == LotEventType.reassigned,
        )
    )
    event_types = [event_type for event_type, _data in published]
    assert "lot.created" in event_types
    assert "lot.renamed" in event_types
    assert "lot.reassigned" in event_types
    assert all(
        "warehouse_id" in data
        for event_type, data in published
        if event_type.startswith("lot.")
    )


def test_staged_receipt_sse_is_warehouse_scoped(client, session, monkeypatch):
    published: list[tuple[str, dict[str, object]]] = []

    async def capture(event_type: str, data: dict[str, object]) -> None:
        published.append((event_type, data))

    monkeypatch.setattr(bus, "publish", capture)
    warehouse = session.get(Warehouse, 1)
    warehouse.receipt_mode = ReceiptMode.admin_review
    session.commit()

    response = client.post(
        "/api/boxes",
        json={
            "box_number": "1",
            "lot": "Staged SSE",
            "pallet_number": "PALLET-STAGED",
            "warehouse_id": 1,
        },
    )

    assert response.status_code == 201
    event = next(data for kind, data in published if kind == "request.created")
    assert event["warehouse_id"] == 1
    assert event["staged"] is True


def test_lot_summary_query_count_is_constant(session, make_user):
    admin = make_user(UserRole.admin)
    for index in range(25):
        create_box(
            session,
            user=admin,
            box_number="1",
            lot=f"Query Lot {index:02d}",
            warehouse_id=1,
        )
    statements = 0

    def count_queries(*_args):
        nonlocal statements
        statements += 1

    event.listen(session.bind, "before_cursor_execute", count_queries)
    try:
        rows, total = list_lot_summaries(
            session,
            user=admin,
            page=1,
            page_size=25,
        )
    finally:
        event.remove(session.bind, "before_cursor_execute", count_queries)
    assert total == len(rows) == 25
    # Three report statements plus at most two ORM refresh/load statements.
    assert statements <= 5
