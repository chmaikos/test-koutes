from __future__ import annotations

import csv
import importlib.util
import io
from pathlib import Path

import pytest
import sqlalchemy as sa
from alembic.config import Config
from alembic.migration import MigrationContext
from alembic.operations import Operations
from alembic.script import ScriptDirectory
from openpyxl import load_workbook
from sqlalchemy import select

from app.deps import get_current_user
from app.main import app
from app.models import (
    Box,
    BoxRequestItem,
    BoxStatus,
    Lot,
    Pallet,
    PalletEvent,
    PalletEventType,
    UserRole,
    Warehouse,
)
from app.models.pallets import clean_pallet_number, normalize_pallet_number
from app.schemas.requests import BoxRequestItemOut, InboundBoxItem
from app.services.requests import merge_inbound_items
from scripts.pallet_preflight import analyze_rows


def _lot(session, name: str = "Pallet Lot") -> Lot:
    lot = Lot(name=name)
    session.add(lot)
    session.commit()
    session.refresh(lot)
    return lot


def _restrict_to(session, user, *warehouse_ids: int) -> None:
    user.warehouses = [
        session.get(Warehouse, warehouse_id) for warehouse_id in warehouse_ids
    ]
    session.commit()


def _load_migration():
    path = (
        Path(__file__).parents[1]
        / "alembic"
        / "versions"
        / "0032_first_class_pallets.py"
    )
    spec = importlib.util.spec_from_file_location("migration_0032", path)
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_pallet_number_cleanup_is_case_preserving_and_lowercased() -> None:
    assert clean_pallet_number("  Pallet\tA-01  ") == "Pallet A-01"
    assert normalize_pallet_number("  Pallet\tA-01  ") == "pallet a-01"


def test_request_item_pallet_snapshot_is_nullable_immutable_and_output_only() -> None:
    legacy = BoxRequestItem(request_id=1, position=1)
    assert legacy.pallet_id is None
    assert legacy.pallet is None

    item = BoxRequestItem(
        id=1,
        request_id=1,
        position=1,
        pallet_id=7,
        pallet="Pallet A-01",
    )
    output = BoxRequestItemOut.model_validate(item)
    assert output.pallet_id == 7
    assert output.pallet == "Pallet A-01"
    assert output.pallet_number == "Pallet A-01"
    assert InboundBoxItem.model_fields["pallet_number"].is_required()
    assert "pallet_id" in InboundBoxItem.model_fields
    assert "pallet" not in InboundBoxItem.model_fields
    with pytest.raises(ValueError, match="pallet snapshot is immutable"):
        item.pallet = "Pallet A-02"


def test_duplicate_inbound_rows_preserve_later_pallet_id() -> None:
    merged = merge_inbound_items(
        [
            InboundBoxItem(
                lot="Lot A",
                box_number="1",
                pallet_number="Pallet A",
                contents="One",
            ),
            InboundBoxItem(
                lot="lot a",
                box_number="001",
                pallet_number=" pallet a ",
                pallet_id=17,
                contents="Two",
            ),
        ]
    )
    assert len(merged) == 1
    assert merged[0].pallet_id == 17
    assert merged[0].contents == "One | Two"


def test_create_scoped_uniqueness_and_legacy_nullable_boxes(
    client, session
) -> None:
    first_lot = _lot(session, "First")
    second_lot = _lot(session, "Second")
    created = client.post(
        "/api/pallets",
        json={
            "lot_id": first_lot.id,
            "warehouse_id": 1,
            "pallet_number": "  Rack\tA-01 ",
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["pallet_number"] == "Rack A-01"
    assert created.json()["normalized_pallet_number"] == "rack a-01"

    collision = client.post(
        "/api/pallets",
        json={
            "lot_id": first_lot.id,
            "warehouse_id": 1,
            "pallet_number": "RACK A-01",
        },
    )
    assert collision.status_code == 409
    assert collision.json()["detail"]["code"] == "pallet_number_collision"

    other_lot = client.post(
        "/api/pallets",
        json={
            "lot_id": second_lot.id,
            "warehouse_id": 1,
            "pallet_number": "RACK A-01",
        },
    )
    assert other_lot.status_code == 201

    legacy = Box(
        box_number="legacy-1",
        lot_id=first_lot.id,
        current_warehouse_id=1,
        status=BoxStatus.received,
    )
    session.add(legacy)
    session.commit()
    assert legacy.pallet_id is None


def test_pallet_summary_csv_and_xlsx_exports(client, session) -> None:
    lot = _lot(session, "Export Lot")
    pallet = Pallet(
        lot_id=lot.id,
        pallet_number="Export Rack",
    )
    session.add(pallet)
    session.flush()
    session.add(
        Box(
            box_number="001",
            lot_id=lot.id,
            pallet_id=pallet.id,
            current_warehouse_id=1,
            status=BoxStatus.received,
        )
    )
    session.add(
        Box(
            box_number="002",
            lot_id=lot.id,
            pallet_id=pallet.id,
            current_warehouse_id=2,
            status=BoxStatus.received,
        )
    )
    session.commit()

    csv_response = client.get(
        "/api/exports/pallets.csv",
        params={"lot_id": lot.id},
    )
    assert csv_response.status_code == 200, csv_response.text
    csv_rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert csv_rows[0][:6] == [
        "Pallet ID",
        "Pallet Number",
        "Lot ID",
        "Lot",
        "Warehouse IDs",
        "Warehouses",
    ]
    assert csv_rows[1][0:4] == [
        str(pallet.id),
        "Export Rack",
        str(lot.id),
        "Export Lot",
    ]
    assert csv_rows[1][4:6] == ["1; 2", "Building 1; Building 2"]
    assert csv_rows[1][7:9] == ["2", "2"]

    xlsx_response = client.get(
        "/api/exports/pallets.xlsx",
        params={"lot_id": lot.id},
    )
    assert xlsx_response.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
    try:
        rows = list(workbook["Pallet summary"].iter_rows(values_only=True))
        assert rows[1][0:4] == (pallet.id, "Export Rack", lot.id, "Export Lot")
        assert rows[1][4:6] == ("1; 2", "Building 1; Building 2")
        assert rows[1][7:9] == (2, 2)
    finally:
        workbook.close()


def test_acl_active_filters_options_and_detail(client, session, make_user) -> None:
    lot = _lot(session)
    warehouse = session.get(Warehouse, 1)
    assert warehouse is not None
    warehouse.name = "North Search Depot"
    one = Pallet(lot_id=lot.id, pallet_number="One")
    two = Pallet(lot_id=lot.id, pallet_number="Two")
    archived = Pallet(
        lot_id=lot.id,
        pallet_number="Old",
        is_active=False,
        archived_at=sa.func.now(),
    )
    session.add_all([one, two, archived])
    session.flush()
    session.add_all(
        [
            Box(
                box_number="one",
                lot_id=lot.id,
                pallet_id=one.id,
                current_warehouse_id=1,
                status=BoxStatus.received,
            ),
            Box(
                box_number="two",
                lot_id=lot.id,
                pallet_id=two.id,
                current_warehouse_id=2,
                status=BoxStatus.received,
            ),
        ]
    )
    session.commit()

    operator = make_user(UserRole.operator)
    _restrict_to(session, operator, 1)
    app.dependency_overrides[get_current_user] = lambda: operator

    listed = client.get("/api/pallets")
    assert listed.status_code == 200
    items = {item["pallet_number"]: item for item in listed.json()["items"]}
    assert set(items) == {"One", "Two"}
    assert items["One"]["warehouse_ids"] == [1]
    assert items["Two"]["warehouse_ids"] == []
    assert items["Two"]["box_count"] == 0
    assert client.get(f"/api/pallets/{two.id}").status_code == 200
    assert client.get("/api/pallets?include_inactive=true").status_code == 403
    options = client.get("/api/pallets/options?search=one")
    assert options.status_code == 200
    assert options.json()["items"][0]["exact_normalized_match"] is True
    warehouse_options = client.get("/api/pallets/options?search=search%20depot")
    assert warehouse_options.status_code == 200
    assert [item["pallet_number"] for item in warehouse_options.json()["items"]] == [
        "One"
    ]
    warehouse_list = client.get("/api/pallets?search=search%20depot")
    assert warehouse_list.status_code == 200
    assert [item["pallet_number"] for item in warehouse_list.json()["items"]] == [
        "One"
    ]


def test_summary_uses_lot_completion_semantics(client, session) -> None:
    lot = _lot(session)
    pallet = Pallet(lot_id=lot.id, pallet_number="Metrics")
    session.add(pallet)
    session.flush()
    for index, status in enumerate(BoxStatus, start=1):
        session.add(
            Box(
                box_number=str(index),
                lot_id=lot.id,
                pallet_id=pallet.id,
                current_warehouse_id=1,
                status=status,
            )
        )
    session.commit()

    response = client.get(f"/api/pallets/{pallet.id}")
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["box_count"] == 6
    assert payload["eligible_box_count"] == 5
    assert payload["completed_box_count"] == 3
    assert payload["completion_percent"] == 60.0
    assert payload["progress_state"] == "in_progress"
    assert payload["status_counts"] == {status.value: 1 for status in BoxStatus}
    assert payload["latest_activity"] is not None


def test_archive_restore_guards_versions_and_events(client, session) -> None:
    lot = _lot(session)
    created = client.post(
        "/api/pallets",
        json={
            "lot_id": lot.id,
            "warehouse_id": 1,
            "pallet_number": "Lifecycle",
        },
    ).json()
    pallet_id = created["id"]
    box = Box(
        box_number="assigned",
        lot_id=lot.id,
        pallet_id=pallet_id,
        current_warehouse_id=1,
        status=BoxStatus.received,
    )
    session.add(box)
    session.commit()

    blocked = client.post(
        f"/api/pallets/{pallet_id}/archive",
        json={"reason": "retire", "expected_version": created["version"]},
    )
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["code"] == "pallet_not_empty"
    assert blocked.json()["detail"]["box_count"] == 1

    box.pallet_id = None
    session.commit()
    archived = client.post(
        f"/api/pallets/{pallet_id}/archive",
        json={"reason": "retire", "expected_version": created["version"]},
    )
    assert archived.status_code == 200, archived.text
    assert archived.json()["is_active"] is False
    assert client.get(f"/api/pallets/{pallet_id}").status_code == 404

    stale = client.post(
        f"/api/pallets/{pallet_id}/restore",
        json={"reason": "return", "expected_version": created["version"]},
    )
    assert stale.status_code == 409
    assert stale.json()["detail"]["code"] == "version_conflict"

    restored = client.post(
        f"/api/pallets/{pallet_id}/restore",
        json={
            "reason": "return",
            "expected_version": archived.json()["version"],
        },
    )
    assert restored.status_code == 200, restored.text
    events = client.get(f"/api/pallets/{pallet_id}/events")
    assert events.status_code == 200
    assert [event["event_type"] for event in reversed(events.json())] == [
        "created",
        "archived",
        "restored",
    ]


def test_rename_requires_admin_and_active_parents(
    client, session, make_user
) -> None:
    lot = _lot(session)
    created = client.post(
        "/api/pallets",
        json={"lot_id": lot.id, "warehouse_id": 1, "pallet_number": "Before"},
    ).json()
    operator = make_user(UserRole.operator)
    app.dependency_overrides[get_current_user] = lambda: operator
    denied = client.patch(
        f"/api/pallets/{created['id']}/rename",
        json={
            "new_pallet_number": "After",
            "reason": "label correction",
            "expected_version": created["version"],
        },
    )
    assert denied.status_code == 403

    admin = make_user(UserRole.admin)
    app.dependency_overrides[get_current_user] = lambda: admin
    warehouse = session.get(Warehouse, 1)
    warehouse.is_active = False
    warehouse.archived_at = sa.func.now()
    session.commit()
    renamed = client.patch(
        f"/api/pallets/{created['id']}/rename",
        json={
            "new_pallet_number": "After",
            "reason": "label correction",
            "expected_version": created["version"],
        },
    )
    assert renamed.status_code == 200


def test_active_pallet_does_not_block_warehouse_archive(client, session) -> None:
    lot = _lot(session)
    session.add(Pallet(lot_id=lot.id, pallet_number="Organizational"))
    session.commit()

    archived = client.delete("/api/warehouses/1")
    assert archived.status_code == 200


def test_preflight_reports_unassigned_and_impossible_assignments() -> None:
    report = analyze_rows(
        [
            {"id": 1, "lot_id": 10, "current_warehouse_id": 1, "pallet_id": None},
            {"id": 2, "lot_id": 11, "current_warehouse_id": 2, "pallet_id": 5},
        ],
        [
            {
                "id": 5,
                "lot_id": 10,
                "normalized_pallet_number": "p-1",
                "is_active": False,
                "merged_into_lot_id": None,
            }
        ],
    )
    assert report["unassigned_box_ids"] == [1]
    assert report["safe"] is False
    assert report["conflicts"]["lot_mismatches"][0]["box_id"] == 2
    assert "warehouse_mismatches" not in report["conflicts"]
    assert report["warehouse_distribution"] == [
        {
            "pallet_id": 5,
            "warehouses": [{"warehouse_id": 2, "box_count": 1}],
        }
    ]
    assert report["conflicts"]["inactive_pallet_assignments"][0]["box_id"] == 2


def test_migration_metadata_offline_sql_and_single_head() -> None:
    assert Pallet.__table__.c.normalized_pallet_number.type.length == 64
    assert "current_warehouse_id" not in Pallet.__table__.c
    assert Box.__table__.c.pallet_id.nullable is True

    migration = _load_migration()
    output = io.StringIO()
    context = MigrationContext.configure(
        dialect_name="postgresql",
        opts={"as_sql": True, "output_buffer": output},
    )
    migration.op = Operations(context)
    migration.upgrade()
    migration.downgrade()
    sql = output.getvalue()
    assert "CREATE TABLE pallets" in sql
    assert "CREATE TABLE pallet_events" in sql
    assert "merged_absorbed" in sql
    assert "lot_reassigned" in sql
    assert "ADD COLUMN pallet_ids JSON" in sql
    assert "ADD COLUMN pallet_count INTEGER" in sql
    assert "ADD COLUMN pallet_snapshots JSON" in sql
    assert "ADD VALUE IF NOT EXISTS 'pallet_assigned'" in sql
    assert "ADD VALUE IF NOT EXISTS 'pallet_unassigned'" in sql
    assert "ADD COLUMN pallet_id INTEGER" in sql
    assert "ALTER TABLE box_request_items ADD COLUMN pallet_id INTEGER" in sql
    assert "ALTER TABLE box_request_items ADD COLUMN pallet VARCHAR(64)" in sql
    assert "FOREIGN KEY(pallet_id) REFERENCES pallets (id) ON DELETE SET NULL" in sql
    assert "CREATE INDEX ix_box_request_items_pallet_id" in sql
    assert "UPDATE boxes" not in sql
    assert "UPDATE box_request_items" not in sql
    assert "INSERT INTO pallets" not in sql
    assert "pallet data, audit snapshots" in sql
    assert "contents over 200 characters exist" in sql

    root = Path(__file__).parents[1]
    config = Config(str(root / "alembic.ini"))
    config.set_main_option("script_location", str(root / "alembic"))
    assert ScriptDirectory.from_config(config).get_heads() == [
        "0033_organizational_pallets"
    ]
    assert migration.down_revision == "0031_return_target_warehouse"


def test_sqlite_migration_preserves_legacy_rows_without_fabricating_data() -> None:
    metadata = sa.MetaData()
    sa.Table("users", metadata, sa.Column("id", sa.Integer(), primary_key=True))
    sa.Table(
        "warehouses",
        metadata,
        sa.Column("id", sa.Integer(), primary_key=True),
    )
    sa.Table(
        "lots",
        metadata,
        sa.Column("id", sa.Integer(), primary_key=True),
    )
    sa.Table(
        "boxes",
        metadata,
        sa.Column("id", sa.Integer(), primary_key=True),
        sa.Column("lot_id", sa.Integer(), nullable=False),
        sa.Column("current_warehouse_id", sa.Integer(), nullable=False),
        sa.Column("contents", sa.String(length=200)),
    )
    sa.Table(
        "box_request_items",
        metadata,
        sa.Column("id", sa.Integer(), primary_key=True),
        sa.Column("contents", sa.String(length=200)),
    )
    sa.Table(
        "lot_purge_events",
        metadata,
        sa.Column("id", sa.Integer(), primary_key=True),
    )
    engine = sa.create_engine("sqlite://")
    metadata.create_all(engine)
    with engine.begin() as connection:
        connection.execute(sa.text("INSERT INTO users (id) VALUES (1)"))
        connection.execute(sa.text("INSERT INTO warehouses (id) VALUES (1)"))
        connection.execute(sa.text("INSERT INTO lots (id) VALUES (1)"))
        connection.execute(
            sa.text(
                "INSERT INTO boxes (id, lot_id, current_warehouse_id, contents) "
                "VALUES (1, 1, 1, 'legacy box contents')"
            )
        )
        connection.execute(
            sa.text(
                "INSERT INTO box_request_items (id, contents) "
                "VALUES (1, 'legacy request contents')"
            )
        )
        migration = _load_migration()
        migration.op = Operations(MigrationContext.configure(connection))
        migration.upgrade()
        assert connection.scalar(sa.text("SELECT COUNT(*) FROM pallets")) == 0
        assert connection.scalar(
            sa.text("SELECT pallet_id FROM boxes WHERE id = 1")
        ) is None
        columns = {
            column["name"]: column
            for column in sa.inspect(connection).get_columns("boxes")
        }
        assert columns["pallet_id"]["nullable"] is True
        assert connection.scalar(
            sa.text("SELECT contents FROM boxes WHERE id = 1")
        ) == "legacy box contents"
        item_columns = {
            column["name"]: column
            for column in sa.inspect(connection).get_columns("box_request_items")
        }
        assert item_columns["pallet_id"]["nullable"] is True
        assert item_columns["pallet"]["nullable"] is True
        assert connection.scalar(
            sa.text("SELECT contents FROM box_request_items WHERE id = 1")
        ) == "legacy request contents"
        assert connection.execute(
            sa.text(
                "SELECT pallet_id, pallet FROM box_request_items WHERE id = 1"
            )
        ).one() == (None, None)
        purge_columns = {
            column["name"]
            for column in sa.inspect(connection).get_columns("lot_purge_events")
        }
        assert {"pallet_ids", "pallet_count", "pallet_snapshots"} <= purge_columns
        item_foreign_keys = sa.inspect(connection).get_foreign_keys(
            "box_request_items"
        )
        assert any(
            foreign_key["constrained_columns"] == ["pallet_id"]
            and foreign_key["referred_table"] == "pallets"
            and foreign_key["options"].get("ondelete") == "SET NULL"
            for foreign_key in item_foreign_keys
        )
        connection.execute(
            sa.text(
                "UPDATE boxes SET contents = :contents WHERE id = 1"
            ),
            {"contents": "x" * 201},
        )
        with pytest.raises(RuntimeError, match="contents over 200 characters"):
            migration.downgrade()
        connection.execute(
            sa.text(
                "UPDATE boxes SET contents = 'legacy box contents' WHERE id = 1"
            )
        )
        connection.execute(
            sa.text(
                "UPDATE lot_purge_events SET pallet_count = 1 WHERE id = 1"
            )
        )
        if (
            connection.scalar(
                sa.text("SELECT COUNT(*) FROM lot_purge_events")
            )
            == 0
        ):
            connection.execute(
                sa.text(
                    "INSERT INTO lot_purge_events "
                    "(id, pallet_count) VALUES (1, 1)"
                )
            )
        with pytest.raises(RuntimeError, match="audit snapshots"):
            migration.downgrade()
        connection.execute(
            sa.text("UPDATE lot_purge_events SET pallet_count = 0")
        )
        migration.downgrade()
        assert "pallet_id" not in {
            column["name"] for column in sa.inspect(connection).get_columns("boxes")
        }
        assert {"pallet_id", "pallet"}.isdisjoint(
            column["name"]
            for column in sa.inspect(connection).get_columns("box_request_items")
        )
        assert {"pallet_ids", "pallet_count", "pallet_snapshots"}.isdisjoint(
            column["name"]
            for column in sa.inspect(connection).get_columns("lot_purge_events")
        )
    engine.dispose()


def test_model_event_metadata_registration(session) -> None:
    assert "pallets" in session.bind.dialect.get_table_names(session.connection())
    assert "pallet_events" in session.bind.dialect.get_table_names(session.connection())
    assert session.scalars(select(PalletEvent)).all() == []
    assert PalletEventType.renumbered.value == "renumbered"
    event_pallet_fk = next(
        foreign_key
        for foreign_key in PalletEvent.__table__.foreign_keys
        if foreign_key.parent.name == "pallet_id"
    )
    assert event_pallet_fk.ondelete == "RESTRICT"
