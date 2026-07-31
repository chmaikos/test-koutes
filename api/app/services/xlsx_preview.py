from __future__ import annotations

import io
from dataclasses import dataclass

from openpyxl import load_workbook

from app.services.boxes import BoxRuleError

MAX_XLSX_BYTES = 5 * 1024 * 1024
MAX_PREVIEW_ROWS = 5000
MAX_PREVIEW_COLUMNS = 50


@dataclass(frozen=True)
class PreviewRow:
    row_number: int
    cells: list[str]


@dataclass(frozen=True)
class PreviewSheet:
    name: str
    max_columns: int
    rows: list[PreviewRow]


def _cell_string(value: object | None) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def preview_xlsx(file_bytes: bytes) -> list[PreviewSheet]:
    if not file_bytes:
        raise BoxRuleError("workbook is empty")
    if len(file_bytes) > MAX_XLSX_BYTES:
        raise BoxRuleError(
            f"file is larger than {MAX_XLSX_BYTES // (1024 * 1024)} MB"
        )
    try:
        workbook = load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise BoxRuleError(f"could not read workbook: {exc}") from exc

    sheets: list[PreviewSheet] = []
    total_rows = 0
    try:
        for worksheet in workbook.worksheets:
            rows: list[PreviewRow] = []
            max_columns = 0
            for row_number, raw_row in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=min(worksheet.max_row or 1, MAX_PREVIEW_ROWS + 1),
                    max_col=min(
                        worksheet.max_column or 1,
                        MAX_PREVIEW_COLUMNS,
                    ),
                    values_only=True,
                ),
                start=1,
            ):
                cells = [_cell_string(value) for value in raw_row]
                while cells and not cells[-1]:
                    cells.pop()
                if not cells or not any(cells):
                    continue
                total_rows += 1
                if total_rows > MAX_PREVIEW_ROWS:
                    raise BoxRuleError(
                        f"workbook exceeds the {MAX_PREVIEW_ROWS}-row preview limit"
                    )
                max_columns = max(max_columns, len(cells))
                rows.append(PreviewRow(row_number=row_number, cells=cells))
            if rows:
                sheets.append(
                    PreviewSheet(
                        name=worksheet.title,
                        max_columns=max_columns,
                        rows=rows,
                    )
                )
    finally:
        workbook.close()

    if not sheets:
        raise BoxRuleError("workbook contains no non-empty rows")
    return sheets
