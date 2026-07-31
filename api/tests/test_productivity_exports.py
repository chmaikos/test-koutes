from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.deps import get_current_user
from app.main import app
from app.models.employees import Employee, ProductivityEntry
from app.models.users import UserRole
from app.models.warehouses import Warehouse


def test_productivity_exports_summary_and_daily_entries(client, session):
    warehouse = session.get(Warehouse, 1)
    warehouse.min_pages_per_day = 100
    employee = Employee(
        warehouse_id=1,
        full_name="Alice",
        default_hours_per_day=Decimal("8"),
    )
    session.add(employee)
    session.flush()
    session.add(
        ProductivityEntry(
            employee_id=employee.id,
            warehouse_id=1,
            entry_date=date(2026, 5, 13),
            pages=50,
            hours_worked=Decimal("8"),
            note="shift",
        )
    )
    session.commit()

    csv_response = client.get(
        "/api/exports/productivity.csv",
        params={"warehouse_id": 1, "report_date": "2026-05-13"},
    )
    assert csv_response.status_code == 200
    assert "Alice" in csv_response.text
    assert "Consistently Below Minimum" in csv_response.text

    xlsx_response = client.get(
        "/api/exports/productivity.xlsx",
        params={"warehouse_id": 1, "report_date": "2026-05-13"},
    )
    assert xlsx_response.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
    try:
        assert workbook.sheetnames == ["Employee averages", "Daily entries"]
        summary_rows = list(workbook["Employee averages"].iter_rows(values_only=True))
        detail_rows = list(workbook["Daily entries"].iter_rows(values_only=True))
        assert summary_rows[1][3] == "Alice"
        assert detail_rows[1][3] == "Alice"
        assert detail_rows[1][5] == 50
        assert detail_rows[1][9] == "shift"
    finally:
        workbook.close()


def test_productivity_export_respects_warehouse_acl(
    client, session, make_user
):
    alice = Employee(warehouse_id=1, full_name="Alice")
    bob = Employee(warehouse_id=2, full_name="Bob")
    session.add_all([alice, bob])
    session.flush()
    session.add_all(
        [
            ProductivityEntry(
                employee_id=alice.id,
                warehouse_id=1,
                entry_date=date(2026, 5, 13),
                pages=50,
                hours_worked=Decimal("8"),
            ),
            ProductivityEntry(
                employee_id=bob.id,
                warehouse_id=2,
                entry_date=date(2026, 5, 13),
                pages=60,
                hours_worked=Decimal("8"),
            ),
        ]
    )
    operator = make_user(UserRole.operator)
    operator.warehouses = [session.get(Warehouse, 1)]
    session.commit()
    app.dependency_overrides[get_current_user] = lambda: operator

    response = client.get(
        "/api/exports/productivity.csv",
        params={"report_date": "2026-05-13"},
    )
    assert response.status_code == 200
    assert "Alice" in response.text
    assert "Bob" not in response.text
