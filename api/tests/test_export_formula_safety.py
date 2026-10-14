"""Every public export family must write untrusted text as spreadsheet text."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from types import SimpleNamespace as NS

import pytest
from openpyxl import load_workbook

from app.models.boxes import BoxStatus
from app.models.requests import BoxRequestStatus
from app.services import exports


@pytest.mark.parametrize("text", ["=1+1", "+cmd", "-cmd", "@SUM(1)", "\t=1+1", " =1+1"])
@pytest.mark.parametrize(
    "family", ["boxes", "files", "lots", "pallets", "productivity", "reconciliation", "analytics"]
)
def test_all_export_families_neutralize_untrusted_formula_text(family, text):
    now = datetime(2026, 10, 14, 12)
    summary = NS(
        id=1,
        barcode="barcode",
        name=text,
        pallet_number=text,
        lot_id=1,
        lot_name=text,
        warehouse_ids=[1],
        warehouse_names=[text],
        is_active=True,
        physical_box_count=1,
        box_count=1,
        active_file_count=1,
        archived_file_count=0,
        status_counts={status.value: 0 for status in BoxStatus},
        eligible_box_count=1,
        completed_box_count=0,
        completion_percent=0,
        progress_state="active",
        warehouse_count=1,
        staged_receipt_count=0,
        last_box_activity=now,
        latest_activity=now,
        created_at=now,
        updated_at=now,
        scope_label=text,
    )
    if family == "boxes":
        box = NS(
            box_number="123",
            barcode="barcode",
            lot=text,
            contents=text,
            current_warehouse_id=1,
            status=BoxStatus.received,
            received_at=now,
            returned_at=None,
            created_at=now,
            updated_at=now,
        )
        csv_data = exports.boxes_to_csv([box], {1: text})
        xlsx_data = exports.boxes_to_xlsx([box], {1: text})
    elif family == "files":
        file = dict(
            position=1,
            status=BoxStatus.received,
            id=1,
            reference=text,
            description=text,
            barcode="barcode",
            lot_id=1,
            lot=text,
            pallet_id=1,
            pallet=text,
            box_id=1,
            box="123",
            warehouse_id=1,
            warehouse=text,
            is_active=True,
            archived_at=None,
            archive_reason=text,
            created_at=now,
            updated_at=now,
        )
        csv_data = exports.files_to_csv([file])
        xlsx_data = exports.files_to_xlsx([file])
    elif family in {"lots", "pallets"}:
        stem = "lot" if family == "lots" else "pallet"
        csv_data = getattr(exports, f"{stem}_summaries_to_csv")([summary])
        xlsx_data = getattr(exports, f"{stem}_summaries_to_xlsx")([summary])
    elif family == "productivity":
        period = NS(
            period_start=date(2026, 10, 12), period_end=date(2026, 10, 18), pages_per_day=100
        )
        average = NS(
            employee_id=1,
            employee_name=text,
            weekly=period,
            monthly=period,
            three_month=period,
            consistently_below_minimum=False,
            excluded_from_metrics=False,
        )
        summary_row = exports.ProductivitySummaryExportRow(1, text, 100, average)
        entry = NS(
            hours_worked=Decimal("8"),
            pages=100,
            entry_date=date(2026, 10, 14),
            excluded_from_metrics=False,
            note=text,
        )
        detail = exports.ProductivityDetailExportRow(1, text, 1, text, entry)
        csv_data = exports.productivity_to_csv([summary_row])
        xlsx_data = exports.productivity_to_xlsx([summary_row], [detail])
    elif family == "reconciliation":
        issue = NS(
            severity="high",
            issue_type="overdue",
            request_id=1,
            box_id=1,
            pallet_id=1,
            pallet_number=text,
            warehouse_id=1,
            warehouse_name=text,
            assigned_mover_user_id=1,
            assigned_mover_name=text,
            request_status=BoxRequestStatus.draft,
            title=text,
            detail=text,
            occurred_at=now,
            due_at=now,
            request_path="/requests/1",
            box_path="/boxes/1",
        )
        csv_data = exports.request_reconciliation_to_csv([issue])
        xlsx_data = exports.request_reconciliation_to_xlsx([issue])
    else:
        duration = NS(supported=True, average_seconds=12, sample_size=1)
        rate = NS(numerator=1, denominator=1, rate=1)
        report = NS(
            total_requests=1,
            completed_requests=1,
            approval_duration=duration,
            preparation_duration=duration,
            transport_duration=duration,
            acceptance_duration=duration,
            on_time=rate,
            discrepancy=rate,
            shortage=rate,
            overage=rate,
            rejection_reasons=[NS(reason=text, count=1)],
            cancellation_reasons=[NS(reason=text, count=1)],
            throughput_by_warehouse=[
                NS(name=text, id=1, completed_requests=1, completed_quantity=1)
            ],
            throughput_by_mover=[],
        )
        csv_data = exports.request_analytics_to_csv(report)
        xlsx_data = exports.request_analytics_to_xlsx(report)

    expected = "'" + text
    csv_rows = list(csv.reader(io.StringIO(csv_data.decode("utf-8-sig"))))
    assert any(expected in row for row in csv_rows[1:])
    assert all(
        not value.lstrip().startswith(("=", "+", "-", "@")) for row in csv_rows[1:] for value in row
    )
    workbook = load_workbook(io.BytesIO(xlsx_data))
    cells = [cell for sheet in workbook for row in sheet.iter_rows(min_row=2) for cell in row]
    assert any(cell.value == expected for cell in cells)
    assert all(cell.data_type != "f" for cell in cells)
    assert any(cell.data_type == "n" and cell.value == 1 for cell in cells)
    if family != "analytics":
        assert any(cell.data_type == "d" for cell in cells)
    workbook.close()
