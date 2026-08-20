from __future__ import annotations

import csv
from datetime import UTC, datetime
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.models.boxes import Box, BoxStatus
from app.models.lots import Lot
from app.models.pallets import Pallet, PalletEvent
from app.models.requests import (
    BoxRequestEvent,
    BoxRequestEventType,
    BoxRequestOrigin,
)
from app.models.users import UserRole
from app.models.warehouses import ReceiptMode, Warehouse
from app.services.boxes import BoxRuleError, create_box
from app.services.requests import create_completed_receipt


def _manual_payload(
    number: str,
    *,
    lot: str = "PALLET-RECEIPT",
    pallet: str = "PALLET-A",
    warehouse_id: int = 1,
    contents: str | None = None,
) -> dict[str, object]:
    return {
        "box_number": number,
        "lot": lot,
        "pallet_number": pallet,
        "warehouse_id": warehouse_id,
        "contents": contents,
    }


def test_manual_receipt_supports_unassigned_and_reuses_explicit_pallet(
    client, session
) -> None:
    pallet_event_count = len(session.scalars(select(PalletEvent)).all())
    missing = client.post(
        "/api/boxes",
        json={"box_number": "1", "lot": "PALLET-RECEIPT", "warehouse_id": 1},
    )
    assert missing.status_code == 201, missing.text
    assert missing.json()["pallet_id"] is None
    assert missing.json()["pallet_number"] is None
    assert len(session.scalars(select(Pallet)).all()) == 0
    assert len(session.scalars(select(PalletEvent)).all()) == pallet_event_count
    unassigned_receipt = client.get(
        f"/api/requests/{missing.json()['receipt_request_id']}"
    ).json()
    assert unassigned_receipt["items"][0]["pallet_id"] is None
    assert unassigned_receipt["items"][0]["pallet"] is None

    first = client.post("/api/boxes", json=_manual_payload("2"))
    second = client.post("/api/boxes", json=_manual_payload("3", pallet=" pallet-a "))
    assert first.status_code == second.status_code == 201
    assert first.json()["pallet_id"] == second.json()["pallet_id"]
    assert first.json()["pallet_number"] == "PALLET-A"
    assert len(session.scalars(select(Pallet)).all()) == 1
    receipt = client.get(f"/api/requests/{first.json()['receipt_request_id']}").json()
    assert receipt["items"][0]["pallet_id"] == first.json()["pallet_id"]
    assert receipt["items"][0]["pallet"] == "PALLET-A"
    sources = client.get(
        "/api/requests/return-sources", params={"warehouse_id": 1}
    ).json()
    source = next(item for item in sources if item["id"] == receipt["id"])
    assert source["pallets"] == [
        {"pallet_id": first.json()["pallet_id"], "pallet_number": "PALLET-A"}
    ]
    assert source["has_unassigned_boxes"] is False


def test_receipt_schemas_canonicalize_whitespace_before_length_validation(
    client,
) -> None:
    manual = client.post(
        "/api/boxes",
        json={
            "box_number": "1",
            "lot": "LONG-BLANK-PALLET",
            "pallet_number": " " * 100,
            "warehouse_id": 1,
        },
    )
    assert manual.status_code == 201, manual.text
    assert (manual.json()["pallet_id"], manual.json()["pallet_number"]) == (
        None,
        None,
    )

    mapped = client.post(
        "/api/boxes/import-mapped",
        json={
            "warehouse_id": 1,
            "items": [
                {
                    "box_number": "2",
                    "lot": "LONG-BLANK-PALLET",
                    "pallet_number": "\t" * 100,
                }
            ],
        },
    )
    assert mapped.status_code == 200, mapped.text
    assert mapped.json()["created"][0]["pallet_id"] is None


def test_box_service_canonicalizes_blank_and_rejects_id_without_number(
    session, make_user
) -> None:
    user = make_user(UserRole.operator)
    unassigned = create_box(
        session,
        user=user,
        box_number="1",
        lot="SERVICE-OPTIONAL",
        warehouse_id=1,
        pallet_number=" \t ",
    )
    assert unassigned.pallet_id is None

    with pytest.raises(BoxRuleError, match="pallet_id requires pallet_number"):
        create_box(
            session,
            user=user,
            box_number="2",
            lot="SERVICE-OPTIONAL",
            warehouse_id=1,
            pallet_id=999,
        )


def test_same_pallet_number_is_reused_in_another_warehouse(client) -> None:
    first = client.post("/api/boxes", json=_manual_payload("1"))
    second = client.post(
        "/api/boxes",
        json=_manual_payload("2", warehouse_id=2),
    )
    assert first.status_code == second.status_code == 201
    assert first.json()["pallet_id"] == second.json()["pallet_id"]


def test_mapped_import_merges_contents_and_rejects_pallet_conflict(client) -> None:
    merged = client.post(
        "/api/boxes/import-mapped",
        json={
            "warehouse_id": 1,
            "items": [
                {
                    "box_number": "1",
                    "lot": "MAPPED-PALLET",
                    "pallet_number": "PALLET-M",
                    "contents": "Invoices",
                },
                {
                    "box_number": "001",
                    "lot": "mapped-pallet",
                    "pallet_number": " pallet-m ",
                    "contents": "Contracts",
                },
            ],
        },
    )
    assert merged.status_code == 200, merged.text
    assert len(merged.json()["created"]) == 1
    assert merged.json()["created"][0]["contents"] == "Contracts | Invoices"

    conflict = client.post(
        "/api/boxes/import-mapped",
        json={
            "warehouse_id": 1,
            "items": [
                {
                    "box_number": "2",
                    "lot": "MAPPED-PALLET",
                    "pallet_number": "PALLET-X",
                },
                {
                    "box_number": "002",
                    "lot": "MAPPED-PALLET",
                    "pallet_number": "PALLET-Y",
                },
            ],
        },
    )
    assert conflict.status_code == 400
    assert "different pallet values" in conflict.json()["detail"]


def test_mapped_import_supports_mixed_assigned_and_unassigned(client) -> None:
    response = client.post(
        "/api/boxes/import-mapped",
        json={
            "warehouse_id": 1,
            "items": [
                {
                    "box_number": "1",
                    "lot": "MAPPED-MIXED",
                    "pallet_number": "PALLET-M",
                },
                {
                    "box_number": "2",
                    "lot": "MAPPED-MIXED",
                },
            ],
        },
    )

    assert response.status_code == 200, response.text
    created = {box["box_number"]: box for box in response.json()["created"]}
    assert created["001"]["pallet_id"] is not None
    assert created["002"]["pallet_id"] is None
    receipt = client.get(
        f"/api/requests/{response.json()['receipt_request_ids'][0]}"
    ).json()
    snapshots = {item["box_number"]: item for item in receipt["items"]}
    assert snapshots["001"]["pallet_id"] == created["001"]["pallet_id"]
    assert snapshots["002"]["pallet_id"] is None
    assert snapshots["002"]["pallet"] is None


def test_raw_xlsx_import_without_pallet_column_creates_unassigned(client) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["box_number", "lot", "warehouse_id"])
    sheet.append(["1", "RAW-PALLET", 1])
    payload = BytesIO()
    workbook.save(payload)
    response = client.post(
        "/api/boxes/import",
        files={
            "file": (
                "missing-pallet.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    assert response.json()["skipped"] == []
    created = response.json()["created"]
    assert len(created) == 1
    assert created[0]["pallet_id"] is None
    assert created[0]["pallet_number"] is None
    receipt = client.get(
        f"/api/requests/{response.json()['receipt_request_ids'][0]}"
    ).json()
    assert receipt["items"][0]["pallet_id"] is None
    assert receipt["items"][0]["pallet"] is None


def test_raw_xlsx_import_supports_mixed_blank_pallet_cells(client) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["box_number", "lot", "pallet_number", "warehouse_id"])
    sheet.append(["1", "RAW-MIXED", "PALLET-R", 1])
    sheet.append(["2", "RAW-MIXED", "   ", 1])
    payload = BytesIO()
    workbook.save(payload)

    response = client.post(
        "/api/boxes/import",
        files={
            "file": (
                "mixed-pallets.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200, response.text
    created = {box["box_number"]: box for box in response.json()["created"]}
    assert created["001"]["pallet_number"] == "PALLET-R"
    assert created["002"]["pallet_id"] is None
    assert created["002"]["pallet_number"] is None


def test_raw_xlsx_rejects_pallet_id_with_whitespace_only_number(client) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        ["box_number", "lot", "pallet_number", "pallet_id", "warehouse_id"]
    )
    sheet.append(["1", "RAW-PALLET-ID", "   ", 42, 1])
    payload = BytesIO()
    workbook.save(payload)

    response = client.post(
        "/api/boxes/import",
        files={
            "file": (
                "blank-pallet-number.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200, response.text
    assert response.json()["created"] == []
    assert response.json()["skipped"][0]["reason"] == (
        "pallet_id requires pallet_number"
    )


def test_raw_xlsx_duplicate_pallet_conflict_blocks_every_duplicate(client) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["box_number", "lot", "pallet_number", "warehouse_id"])
    sheet.append(["1", "RAW-CONFLICT", "PALLET-A", 1])
    sheet.append(["001", "raw-conflict", "   ", 1])
    sheet.append(["1", "RAW-CONFLICT", "PALLET-A", 1])
    payload = BytesIO()
    workbook.save(payload)

    response = client.post(
        "/api/boxes/import",
        files={
            "file": (
                "conflicting-pallets.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 200, response.text
    assert response.json()["created"] == []
    assert [entry["row"] for entry in response.json()["skipped"]] == [2, 3, 4]
    assert all(
        "Assigned vs Unassigned" in entry["reason"]
        for entry in response.json()["skipped"]
    )


def test_staged_snapshot_survives_pallet_rename_and_finalizes(client, session) -> None:
    warehouse = session.get(Warehouse, 1)
    assert warehouse is not None
    warehouse.receipt_mode = ReceiptMode.admin_review
    session.commit()

    staged = client.post(
        "/api/boxes",
        json=_manual_payload("10", lot="STAGED-PALLET", pallet="PALLET-BEFORE"),
    )
    request_id = staged.json()["staged_receipt_id"]
    request = client.get(f"/api/requests/{request_id}").json()
    item = request["items"][0]
    pallet = client.get(f"/api/pallets/{item['pallet_id']}").json()
    renamed = client.patch(
        f"/api/pallets/{item['pallet_id']}/rename",
        json={
            "new_pallet_number": "PALLET-AFTER",
            "reason": "Correct label",
            "expected_version": pallet["version"],
        },
    )
    assert renamed.status_code == 200, renamed.text

    finalized = client.post(
        f"/api/requests/{request_id}/approve",
        json={"expected_version": request["version"]},
    )
    assert finalized.status_code == 200, finalized.text
    assert finalized.json()["items"][0]["pallet"] == "PALLET-BEFORE"
    box = session.get(Box, finalized.json()["items"][0]["box_id"])
    assert box is not None
    assert box.pallet_id == item["pallet_id"]
    assert box.pallet.pallet_number == "PALLET-AFTER"


def test_manual_unassigned_receipt_stages_and_finalizes(client, session) -> None:
    warehouse = session.get(Warehouse, 1)
    assert warehouse is not None
    warehouse.receipt_mode = ReceiptMode.admin_review
    session.commit()

    staged = client.post(
        "/api/boxes",
        json={
            "box_number": "1",
            "lot": "STAGED-UNASSIGNED",
            "warehouse_id": 1,
        },
    )
    assert staged.status_code == 201, staged.text
    request_id = staged.json()["staged_receipt_id"]
    request = client.get(f"/api/requests/{request_id}").json()
    assert (request["items"][0]["pallet_id"], request["items"][0]["pallet"]) == (
        None,
        None,
    )

    finalized = client.post(
        f"/api/requests/{request_id}/approve",
        json={"expected_version": request["version"]},
    )
    assert finalized.status_code == 200, finalized.text
    box = session.get(Box, finalized.json()["items"][0]["box_id"])
    assert box is not None and box.pallet_id is None


def test_mixed_staged_receipt_finalizes_assigned_and_unassigned(
    client, session
) -> None:
    warehouse = session.get(Warehouse, 1)
    assert warehouse is not None
    warehouse.receipt_mode = ReceiptMode.admin_review
    session.commit()

    staged = client.post(
        "/api/boxes/import-mapped",
        json={
            "warehouse_id": 1,
            "items": [
                {
                    "box_number": "1",
                    "lot": "STAGED-MIXED",
                    "pallet_number": "PALLET-S",
                },
                {
                    "box_number": "2",
                    "lot": "STAGED-MIXED",
                },
            ],
        },
    )
    assert staged.status_code == 200, staged.text
    assert staged.json()["created"] == []
    assert len(staged.json()["staged_receipt_ids"]) == 1
    request_id = staged.json()["staged_receipt_ids"][0]
    request = client.get(f"/api/requests/{request_id}").json()
    snapshots = {item["box_number"]: item for item in request["items"]}
    assert snapshots["001"]["pallet_id"] is not None
    assert (snapshots["002"]["pallet_id"], snapshots["002"]["pallet"]) == (
        None,
        None,
    )

    finalized = client.post(
        f"/api/requests/{request_id}/approve",
        json={"expected_version": request["version"]},
    )
    assert finalized.status_code == 200, finalized.text
    boxes = {
        item["box_number"]: session.get(Box, item["box_id"])
        for item in finalized.json()["items"]
    }
    assert boxes["001"] is not None and boxes["001"].pallet_id is not None
    assert boxes["002"] is not None and boxes["002"].pallet_id is None


def test_legacy_unassigned_box_remains_returnable_and_snapshots_null(
    client, session, make_user
) -> None:
    operator = make_user(UserRole.operator)
    lot = Lot(name="LEGACY-RETURN")
    session.add(lot)
    session.flush()
    box = Box(
        box_number="001",
        lot_id=lot.id,
        current_warehouse_id=1,
        status=BoxStatus.ready_to_return,
        received_at=datetime.now(UTC),
    )
    session.add(box)
    session.flush()
    source = create_completed_receipt(
        session,
        user=operator,
        warehouse_id=1,
        boxes=[box],
        origin=BoxRequestOrigin.manual_entry,
        note="Legacy receipt compatibility.",
    )
    assert source.items[0].pallet_id is None
    assert source.items[0].pallet is None
    sources = client.get(
        "/api/requests/return-sources", params={"warehouse_id": 1}
    ).json()
    legacy_source = next(item for item in sources if item["id"] == source.id)
    assert legacy_source["pallets"] == []
    assert legacy_source["has_unassigned_boxes"] is True

    candidates = client.get(f"/api/requests/{source.id}/return-candidates")
    assert candidates.status_code == 200
    assert candidates.json()[0]["pallet_id"] is None
    assert candidates.json()[0]["pallet_number"] is None
    created = client.post(
        "/api/requests",
        json={
            "direction": "return",
            "warehouse_id": 1,
            "target_warehouse_id": 2,
            "quantity": 1,
            "source_inbound_request_id": source.id,
            "box_ids": [box.id],
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["items"][0]["pallet_id"] is None
    assert created.json()["items"][0]["pallet"] is None


def test_unassigned_exports_are_blank_and_xlsx_reimports(client) -> None:
    created = client.post(
        "/api/boxes",
        json={
            "box_number": "1",
            "lot": "EXPORT-UNASSIGNED",
            "warehouse_id": 1,
        },
    )
    assert created.status_code == 201, created.text

    csv_response = client.get("/api/exports/boxes.csv")
    rows = list(
        csv.DictReader(csv_response.content.decode("utf-8-sig").splitlines())
    )
    exported = next(row for row in rows if row["Box Number"] == "001")
    assert exported["Pallet ID"] == ""
    assert exported["Pallet Number"] == ""

    xlsx_response = client.get("/api/exports/boxes.xlsx")
    workbook = load_workbook(BytesIO(xlsx_response.content))
    sheet = workbook.active
    assert sheet is not None
    headers = [cell.value for cell in sheet[1]]
    box_column = headers.index("Box Number") + 1
    lot_column = headers.index("Lot") + 1
    pallet_id_column = headers.index("Pallet ID") + 1
    pallet_number_column = headers.index("Pallet Number") + 1
    assert sheet.cell(2, pallet_id_column).value is None
    assert sheet.cell(2, pallet_number_column).value is None
    sheet.cell(2, box_column, "2")
    sheet.cell(2, lot_column, "EXPORT-ROUNDTRIP")
    payload = BytesIO()
    workbook.save(payload)

    imported = client.post(
        "/api/boxes/import",
        files={
            "file": (
                "roundtrip.xlsx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert imported.status_code == 200, imported.text
    assert imported.json()["skipped"] == []
    assert imported.json()["created"][0]["pallet_id"] is None
    assert imported.json()["created"][0]["pallet_number"] is None


def test_request_reporting_search_export_and_long_contents_include_pallet(
    client, session
) -> None:
    contents = "x" * 1500
    created = client.post(
        "/api/boxes",
        json=_manual_payload(
            "20",
            lot="REPORT-PALLET",
            pallet="SEARCHABLE-PALLET",
            contents=contents,
        ),
    )
    assert created.status_code == 201, created.text
    assert created.json()["contents"] == contents
    assert client.get("/api/boxes", params={"search": "searchable-pallet"}).json()[
        "total"
    ] == 1
    requests = client.get(
        "/api/requests", params={"search": "searchable-pallet"}
    ).json()
    assert requests["total"] == 1
    assert requests["items"][0]["items"][0]["pallet_number"] == "SEARCHABLE-PALLET"
    exported = client.get("/api/exports/boxes.csv")
    assert exported.status_code == 200
    header = exported.content.decode("utf-8-sig").splitlines()[0]
    assert "Pallet ID" in header and "Pallet Number" in header

    receipt_id = created.json()["receipt_request_id"]
    session.add(
        BoxRequestEvent(
            request_id=receipt_id,
            event_type=BoxRequestEventType.cancelled,
            user_id=None,
            note="Reporting context",
            event_metadata={
                "admin_override": True,
                "cancelled_reservation": True,
                "box_id": created.json()["id"],
            },
        )
    )
    session.commit()
    report = client.get("/api/requests/reconciliation").json()
    issue = next(
        item
        for item in report["items"]
        if item["request_id"] == receipt_id
        and item["issue_type"] == "cancelled_reservation"
    )
    assert issue["pallet_id"] == created.json()["pallet_id"]
    assert issue["pallet_number"] == "SEARCHABLE-PALLET"
