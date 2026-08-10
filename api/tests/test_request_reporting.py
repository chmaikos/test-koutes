from __future__ import annotations

import io
from datetime import UTC, datetime, timedelta

from openpyxl import load_workbook
from sqlalchemy import event, select

from app.deps import get_current_user
from app.main import app
from app.models.requests import (
    BoxRequest,
    BoxRequestDirection,
    BoxRequestDiscrepancy,
    BoxRequestDiscrepancyType,
    BoxRequestDocument,
    BoxRequestDocumentType,
    BoxRequestEvent,
    BoxRequestEventType,
    BoxRequestOrigin,
    BoxRequestPriority,
    BoxRequestStatus,
)
from app.models.users import UserRole
from app.models.warehouses import Warehouse
from app.services.request_reporting import (
    RequestReportFilters,
    build_reconciliation,
)


def _request(
    *,
    warehouse_id: int,
    requester_id: int,
    status: BoxRequestStatus,
    submitted_at: datetime,
    quantity: int = 10,
    actual: int | None = None,
    variance: int | None = None,
) -> BoxRequest:
    return BoxRequest(
        direction=BoxRequestDirection.inbound,
        warehouse_id=warehouse_id,
        quantity=quantity,
        status=status,
        requester_user_id=requester_id,
        priority=BoxRequestPriority.normal,
        origin=BoxRequestOrigin.workflow,
        suggestion_quantity=quantity,
        current_available=0,
        min_inventory=0,
        pending_inbound=0,
        eligible_return=0,
        actual_received_quantity=actual,
        variance_quantity=variance,
        submitted_at=submitted_at,
        created_at=submitted_at,
        updated_at=submitted_at,
        version=1,
    )


def test_reconciliation_filters_acl_dates_and_structured_events(
    client, session, make_user
):
    operator = make_user(UserRole.operator)
    operator.warehouses = [session.get(Warehouse, 1)]
    now = datetime(2026, 8, 10, 12, tzinfo=UTC)
    shortage = _request(
        warehouse_id=1,
        requester_id=operator.id,
        status=BoxRequestStatus.completed,
        submitted_at=now - timedelta(days=2),
        actual=8,
        variance=-2,
    )
    hidden = _request(
        warehouse_id=2,
        requester_id=operator.id,
        status=BoxRequestStatus.in_transit,
        submitted_at=now,
    )
    session.add_all([shortage, hidden])
    session.flush()
    shortage.completed_at = now
    session.add(
        BoxRequestEvent(
            request_id=shortage.id,
            event_type=BoxRequestEventType.cancelled,
            from_status=BoxRequestStatus.approved,
            to_status=BoxRequestStatus.cancelled,
            user_id=operator.id,
            note="Reservation released",
            occurred_at=now,
            event_metadata={
                "cancelled_reservation": True,
                "box_id": 99,
            },
        )
    )
    session.commit()
    app.dependency_overrides[get_current_user] = lambda: operator

    response = client.get(
        "/api/requests/reconciliation",
        params={
            "from_at": (now - timedelta(minutes=1)).isoformat(),
            "to_at": (now + timedelta(minutes=1)).isoformat(),
        },
    )
    assert response.status_code == 200
    payload = response.json()
    assert {item["issue_type"] for item in payload["items"]} == {
        "shortage",
        "missing_erp_document",
        "cancelled_reservation",
    }
    assert all(item["warehouse_id"] == 1 for item in payload["items"])
    assert client.get(
        "/api/requests/reconciliation", params={"severity": "critical"}
    ).json()["total"] == 0


def test_request_analytics_denominators_nulls_and_throughput(
    client, session, make_user
):
    operator = make_user(UserRole.operator)
    operator.warehouses = [session.get(Warehouse, 1)]
    start = datetime(2026, 8, 1, 8, tzinfo=UTC)
    completed = _request(
        warehouse_id=1,
        requester_id=operator.id,
        status=BoxRequestStatus.completed,
        submitted_at=start,
        quantity=5,
        actual=4,
        variance=-1,
    )
    completed.approved_at = start + timedelta(hours=1)
    completed.ready_for_transport_at = start + timedelta(hours=2)
    completed.in_transit_at = start + timedelta(hours=3)
    completed.awaiting_confirmation_at = start + timedelta(hours=4)
    completed.completed_at = start + timedelta(hours=5)
    completed.sla_deadline = start + timedelta(hours=6)
    rejected = _request(
        warehouse_id=1,
        requester_id=operator.id,
        status=BoxRequestStatus.rejected,
        submitted_at=start,
    )
    rejected.rejection_reason = "Capacity"
    session.add_all([completed, rejected])
    session.flush()
    session.add(
        BoxRequestDiscrepancy(
            request_id=completed.id,
            discrepancy_type=BoxRequestDiscrepancyType.missing,
            quantity=1,
            created_at=completed.completed_at,
        )
    )
    session.commit()
    app.dependency_overrides[get_current_user] = lambda: operator

    response = client.get("/api/requests/analytics")
    assert response.status_code == 200
    payload = response.json()
    assert payload["total_requests"] == 2
    assert payload["completed_requests"] == 1
    assert payload["approval_duration"]["average_seconds"] == 3600
    assert payload["preparation_duration"] == {
        "supported": True,
        "average_seconds": 3600,
        "sample_size": 1,
    }
    assert payload["transport_duration"]["average_seconds"] == 3600
    assert payload["acceptance_duration"]["average_seconds"] == 3600
    assert payload["on_time"] == {"numerator": 1, "denominator": 1, "rate": 1.0}
    assert payload["shortage"] == {"numerator": 1, "denominator": 1, "rate": 1.0}
    assert payload["overage"] == {"numerator": 0, "denominator": 1, "rate": 0.0}
    assert payload["rejection_reasons"] == [{"reason": "Capacity", "count": 1}]
    assert payload["throughput_by_warehouse"][0]["completed_quantity"] == 4


def test_request_exports_are_filtered_acl_safe_and_timestamped(
    client, session, make_user
):
    operator = make_user(UserRole.operator)
    operator.warehouses = [session.get(Warehouse, 1)]
    now = datetime.now(UTC)
    visible = _request(
        warehouse_id=1,
        requester_id=operator.id,
        status=BoxRequestStatus.completed,
        submitted_at=now,
        actual=1,
        variance=-1,
    )
    hidden = _request(
        warehouse_id=2,
        requester_id=operator.id,
        status=BoxRequestStatus.completed,
        submitted_at=now,
        actual=1,
        variance=1,
    )
    session.add_all([visible, hidden])
    session.commit()
    app.dependency_overrides[get_current_user] = lambda: operator

    csv_response = client.get("/api/exports/request-reconciliation.csv")
    assert csv_response.status_code == 200
    assert "request-reconciliation-" in csv_response.headers["content-disposition"]
    assert "Building 1" in csv_response.text
    assert "Building 2" not in csv_response.text

    xlsx_response = client.get("/api/exports/request-analytics.xlsx")
    assert xlsx_response.status_code == 200
    assert "request-analytics-" in xlsx_response.headers["content-disposition"]
    workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
    try:
        assert workbook.sheetnames == ["Request analytics"]
    finally:
        workbook.close()


def test_reconciliation_query_count_is_constant(session, make_user):
    admin = make_user(UserRole.admin)
    now = datetime.now(UTC)
    session.add_all(
        [
            _request(
                warehouse_id=1,
                requester_id=admin.id,
                status=BoxRequestStatus.in_transit,
                submitted_at=now,
            )
            for _ in range(25)
        ]
    )
    session.commit()
    statements = 0

    def count_queries(*_args):
        nonlocal statements
        statements += 1

    event.listen(session.bind, "before_cursor_execute", count_queries)
    try:
        report = build_reconciliation(
            session,
            user=admin,
            filters=RequestReportFilters(),
            now=now,
        )
    finally:
        event.remove(session.bind, "before_cursor_execute", count_queries)
    assert report.total >= 25
    # Six reporting queries plus at most two ORM refresh/select-in queries.
    assert statements <= 8


def test_document_upload_metadata_tracks_superseded_ids(session, make_user):
    admin = make_user(UserRole.admin)
    now = datetime.now(UTC)
    request = _request(
        warehouse_id=1,
        requester_id=admin.id,
        status=BoxRequestStatus.approved,
        submitted_at=now,
    )
    session.add(request)
    session.flush()
    old = BoxRequestDocument(
        request_id=request.id,
        document_type=BoxRequestDocumentType.delivery_note,
        erp_reference="OLD",
        object_key="old",
        original_filename="old.pdf",
        content_type="application/pdf",
        size_bytes=1,
        sha256="0" * 64,
        is_current=False,
    )
    session.add(old)
    session.commit()
    event_row = BoxRequestEvent(
        request_id=request.id,
        event_type=BoxRequestEventType.document_uploaded,
        from_status=request.status,
        to_status=request.status,
        user_id=admin.id,
        event_metadata={"superseded_document_ids": [old.id]},
    )
    session.add(event_row)
    session.commit()
    stored = session.scalar(
        select(BoxRequestEvent).where(BoxRequestEvent.id == event_row.id)
    )
    assert stored.event_metadata["superseded_document_ids"] == [old.id]
